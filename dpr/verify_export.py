"""Same column-offset proof, run against a DPR produced by the HTML tool."""
import re, sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1] if len(sys.argv) > 1 else 'test-output.xlsx')
fail = []
bk = wb['Piling']
hdr = {c.column_letter: c.value for c in bk[7] if c.value}
checks = [('O9', 'Variance for the Day',    ['Plan FTD', 'Achieved FTD']),
          ('Q9', 'Achieved FTM till date',  ['Achieved FTD', 'Achieved FTM till previous date']),
          ('R9', 'Workdone Till Date',      ['Achieved till last month', 'Achieved FTM till date']),
          ('S9', '% Complete',              ['Workdone Till Date', 'Scope']),
          ('T9', 'Balance',                 ['Scope', 'Workdone Till Date'])]
for cell, owner, operands in checks:
    if hdr.get(cell[0]) != owner:
        fail.append('%s sits under %r, expected %r' % (cell, hdr.get(cell[0]), owner)); continue
    f = bk[cell].value or ''
    seen = [hdr.get(l) for l in re.findall(r'\$?([A-Z])9', f) if l != cell[0]]
    for need in operands:
        if need not in seen:
            fail.append('%s (%s) references %s, missing %r' % (cell, owner, seen, need))

# subtotals on the summary must add up the detail rows above them, nothing else
sm = wb['Summary-AreaWise']
subs = 0
for r in range(8, sm.max_row + 1):
    lab = sm.cell(row=r, column=3).value
    if not isinstance(lab, str): continue
    if lab.startswith('Sub total') or lab == 'GRAND TOTAL':
        subs += 1
        f = sm.cell(row=r, column=5).value or ''
        refs = [int(x) for x in re.findall(r'E(\d+)', f)]
        if not refs: fail.append('row %d (%s) has no SUM range' % (r, lab)); continue
        if max(refs) >= r: fail.append('row %d (%s) sums a row at or below itself' % (r, lab))
        for rr in refs:
            other = sm.cell(row=rr, column=3).value
            is_detail = isinstance(other, str) and not other.startswith('Sub total') and other != 'GRAND TOTAL'
            is_sub = isinstance(other, str) and other.startswith('Sub total')
            if lab == 'GRAND TOTAL' and not is_sub:
                fail.append('GRAND TOTAL sums row %d which is %r, not a sub total' % (rr, other))
            if lab.startswith('Sub total') and not is_detail:
                fail.append('%s sums row %d which is %r, not a detail row' % (lab, rr, other))
        pct = sm.cell(row=r, column=9).value or ''
        if 'H%d/E%d' % (r, r) not in pct.replace(' ', ''):
            fail.append('row %d %% Complete is %r' % (r, pct))
print('checked %d total rows on the summary' % subs)
if fail:
    print('FAILURES (%d):' % len(fail))
    for f in fail: print('  -', f)
    sys.exit(1)
print('OK — exported workbook is internally consistent.')
