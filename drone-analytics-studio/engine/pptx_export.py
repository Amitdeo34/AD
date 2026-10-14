"""PowerPoint report builder.

Analysis frames come in as images rendered by the studio (identical to the video), everything else
(titles, status tables, concerns, milestones, observations) is native, editable PowerPoint content.
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

FONT = 'Segoe UI'
INK = '0F172A'
MUTED = '64748B'
LINE = 'E2E8F0'
SOFT = 'F8FAFC'
WHITE = 'FFFFFF'
SEVERITY = {'high': 'EF4444', 'medium': 'F59E0B', 'low': '3B82F6'}

SW, SH = Inches(13.333), Inches(7.5)


def _hex(h, default=INK):
    h = (h or default).lstrip('#')
    if len(h) == 3:
        h = ''.join(c * 2 for c in h)
    return h.upper()[:6]


def _rgb(h):
    return RGBColor.from_string(_hex(h))


def _set_alpha(fill_parent_xml, alpha):
    srgb = fill_parent_xml.find('.//' + qn('a:srgbClr'))
    if srgb is not None:
        a = etree.SubElement(srgb, qn('a:alpha'))
        a.set('val', str(int(alpha * 100000)))


def box(slide, x, y, w, h, fill=None, line=None, alpha=None, shape=MSO_SHAPE.RECTANGLE, radius=None, lw=0.75):
    shp = slide.shapes.add_shape(shape, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(fill)
        if alpha is not None:
            _set_alpha(shp.fill._xPr, alpha)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = _rgb(line)
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if shp.has_text_frame:
        shp.text_frame.text = ''
    return shp


def text(slide, x, y, w, h, paras, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=None, margin=0.0):
    """paras: str | list of paragraphs; a paragraph is str or list of runs (text, {opts})."""
    tb = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ('left', 'right', 'top', 'bottom'):
        setattr(tf, 'margin_' + side, Inches(margin))
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = str(t)
            f = r.font
            f.name = o.get('font', FONT)
            f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold)
            f.italic = o.get('italic', False)
            f.color.rgb = _rgb(o.get('color', color))
            if o.get('spc'):
                rPr = r._r.get_or_add_rPr()
                rPr.set('spc', str(int(o['spc'])))
    return tb


def picture_cover(slide, path, x, y, w, h):
    """Insert a picture scaled to cover the box (cropping the overflow) like CSS object-fit: cover."""
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    pic = slide.shapes.add_picture(path, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    box_r, img_r = w / h, iw / ih
    if img_r > box_r:
        crop = (1 - box_r / img_r) / 2
        pic.crop_left = pic.crop_right = crop
    elif img_r < box_r:
        crop = (1 - img_r / box_r) / 2
        pic.crop_top = pic.crop_bottom = crop
    return pic


class Deck:
    def __init__(self, spec, img_dir):
        self.spec = spec
        self.info = spec.get('info', {})
        self.accent = _hex(self.info.get('accent'), 'F59E0B')
        self.img_dir = img_dir
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SW, SH
        self.blank = self.prs.slide_layouts[6]
        self.page = 0

    def img(self, name):
        if not name:
            return None
        p = os.path.join(self.img_dir, os.path.basename(name))
        return p if os.path.exists(p) else None

    def new_slide(self, notes=None):
        s = self.prs.slides.add_slide(self.blank)
        self.page += 1
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    # ------------------------------------------------------------------ chrome
    def header(self, s, title, subtitle=None):
        box(s, 0, 0, SW, Inches(0.09), fill=self.accent)
        text(s, Inches(0.55), Inches(0.32), Inches(9.3), Inches(0.55), title, size=24, bold=True)
        if subtitle:
            text(s, Inches(0.55), Inches(0.86), Inches(9.3), Inches(0.35), subtitle, size=12, color=MUTED)
        right = [[(self.info.get('name', ''), {'bold': True, 'size': 11, 'color': INK})],
                 [(' · '.join(x for x in [self.info.get('surveyLabel'), self.info.get('surveyDateText')] if x),
                   {'size': 10, 'color': MUTED})]]
        text(s, Inches(9.4), Inches(0.34), Inches(3.4), Inches(0.7), right, align=PP_ALIGN.RIGHT)

    def footer(self, s):
        box(s, Inches(0.55), Inches(7.02), SW - Inches(1.1), Emu(9525), fill=LINE)
        left = ' · '.join(x for x in ['Drone Progress Report', self.info.get('name'), self.info.get('company')] if x)
        text(s, Inches(0.55), Inches(7.08), Inches(9), Inches(0.3), left, size=9, color=MUTED)
        text(s, SW - Inches(2.55), Inches(7.08), Inches(2.0), Inches(0.3), str(self.page), size=9, color=MUTED,
             align=PP_ALIGN.RIGHT)

    # ------------------------------------------------------------------ slides
    def title(self, sl):
        s = self.new_slide()
        cover = self.img(self.spec.get('images', {}).get('cover'))
        if cover:
            picture_cover(s, cover, 0, 0, SW, SH)
            box(s, 0, 0, SW, SH, fill='020617', alpha=0.55)
            box(s, 0, 0, Inches(7.8), SH, fill='020617', alpha=0.35)
        else:
            box(s, 0, 0, SW, SH, fill='0B1220')
        box(s, Inches(0.8), Inches(2.25), Inches(0.09), Inches(2.6), fill=self.accent)
        text(s, Inches(1.1), Inches(2.15), Inches(10), Inches(0.4),
             [[(sl.get('overline') or 'DRONE PROGRESS REPORT', {'bold': True, 'size': 13, 'color': self.accent, 'spc': 300})]])
        text(s, Inches(1.1), Inches(2.6), Inches(11), Inches(1.3), sl.get('title') or self.info.get('name', 'Project'),
             size=44, bold=True, color=WHITE)
        if self.info.get('site'):
            text(s, Inches(1.1), Inches(3.75), Inches(11), Inches(0.5), self.info['site'], size=18, color='CBD5E1')
        meta = [x for x in [self.info.get('surveyLabel'), self.info.get('surveyDateText'),
                            ('Client: ' + self.info['client']) if self.info.get('client') else None] if x]
        if meta:
            text(s, Inches(1.1), Inches(4.35), Inches(11), Inches(0.45), '   |   '.join(meta), size=14, color='E2E8F0')
        by = ' · '.join(x for x in [('Prepared by ' + self.info['preparedBy']) if self.info.get('preparedBy') else None,
                                     self.info.get('company')] if x)
        if by:
            text(s, Inches(1.1), Inches(6.55), Inches(11), Inches(0.4), by, size=12, color='94A3B8')
        logo = self.img(self.spec.get('images', {}).get('logo'))
        if logo:
            from PIL import Image
            with Image.open(logo) as im:
                r = im.width / max(1, im.height)
            hh = Inches(0.7)
            s.shapes.add_picture(logo, Inches(1.1), Inches(0.75), int(hh * r), hh)

    def analysis(self, sl):
        lines = []
        for z in sl.get('zones', []):
            lines.append(f"{z.get('name')}: {z.get('stage')} {z.get('progress', '')}%" + (' (CRITICAL)' if z.get('critical') else ''))
        s = self.new_slide(notes='\n'.join(lines + sl.get('notes', [])))
        no = sl.get('no')
        self.header(s, (f'Analysis {no:02d} — ' if no else '') + (sl.get('title') or 'Site view'),
                    sl.get('caption') or None)
        x, y, w = Inches(0.55), Inches(1.32), Inches(8.65)
        h = int(w * 9 / 16)
        path = self.img(sl.get('image'))
        if path:
            picture_cover(s, path, x, y, w, h)
        box(s, x, y, w, h, line=LINE)
        if sl.get('source'):
            text(s, x, y + h + Inches(0.06), w, Inches(0.3), sl['source'], size=9, color=MUTED)

        px, pw = Inches(9.45), Inches(3.33)
        box(s, px, y, pw, h, fill=SOFT, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
        paras = []
        zones = sl.get('zones', [])
        if zones:
            paras.append([('BUILDINGS / ZONES', {'bold': True, 'size': 9, 'color': MUTED, 'spc': 120})])
            for z in zones:
                paras.append([('■ ', {'color': z.get('color', MUTED), 'size': 12}),
                              (z.get('name', ''), {'bold': True, 'size': 12})] +
                             ([('   CRITICAL', {'bold': True, 'size': 9, 'color': 'DC2626'})] if z.get('critical') else []))
                det = ' · '.join(str(v) for v in [z.get('stage'), f"{z['progress']}%" if z.get('progress') is not None else None,
                                                  z.get('floors')] if v)
                paras.append([('     ' + det, {'size': 10, 'color': '334155'})])
                if z.get('remarks'):
                    paras.append([('     ' + z['remarks'], {'size': 9, 'color': MUTED, 'italic': True})])
        pins = sl.get('pins', [])
        if pins:
            paras.append([(' ', {'size': 6})])
            paras.append([('CONCERNS', {'bold': True, 'size': 9, 'color': MUTED, 'spc': 120})])
            for p in pins:
                col = SEVERITY.get(p.get('severity'), 'EF4444')
                paras.append([(f"● {p.get('label', '')} ", {'bold': True, 'size': 11, 'color': col}),
                              (p.get('title', ''), {'size': 11, 'bold': True})])
                if p.get('detail'):
                    paras.append([('     ' + p['detail'], {'size': 9, 'color': '475569'})])
        notes = sl.get('notes', [])
        if notes:
            paras.append([(' ', {'size': 6})])
            paras.append([('OBSERVATIONS', {'bold': True, 'size': 9, 'color': MUTED, 'spc': 120})])
            for n in notes:
                paras.append([('•  ' + n, {'size': 10, 'color': '1E293B'})])
        if not paras:
            paras = [[('Add buildings, concern pins or observations to this analysis in the studio.',
                       {'size': 10, 'color': MUTED, 'italic': True})]]
        text(s, px + Inches(0.18), y + Inches(0.16), pw - Inches(0.36), h - Inches(0.3), paras, spacing=2)
        self.footer(s)

    def compare(self, sl):
        s = self.new_slide()
        self.header(s, sl.get('title') or 'Progress comparison', sl.get('caption') or None)
        w = Inches(6.0)
        h = int(w * 9 / 16)
        y = Inches(1.75)
        for i, (img, label) in enumerate([(sl.get('imageA'), sl.get('labelA')), (sl.get('imageB'), sl.get('labelB'))]):
            x = Inches(0.55) + i * (w + Inches(0.23))
            p = self.img(img)
            if p:
                picture_cover(s, p, x, y, w, h)
            box(s, x, y, w, h, line=LINE)
            chip = box(s, x + Inches(0.15), y + Inches(0.15), Inches(1.9), Inches(0.38),
                       fill=self.accent if i else '334155', shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
            tf = chip.text_frame
            tf.text = (label or ('Before' if i == 0 else 'After')).upper()
            for r in tf.paragraphs[0].runs:
                r.font.size, r.font.bold, r.font.name = Pt(11), True, FONT
                r.font.color.rgb = _rgb(WHITE)
        self.footer(s)

    def summary(self, sl):
        s = self.new_slide()
        page = f" ({sl['page']}/{sl['pages']})" if sl.get('pages', 1) > 1 else ''
        self.header(s, (sl.get('title') or 'Building Status Summary') + page,
                    'Construction stage and progress of each building as observed in this survey')
        kpis = sl.get('kpis', [])
        kx, ky = Inches(0.55), Inches(1.4)
        kw = (SW - Inches(1.1) - Inches(0.2) * (len(kpis) - 1)) / max(1, len(kpis)) if kpis else 0
        for i, k in enumerate(kpis):
            x = kx + i * (kw + Inches(0.2))
            box(s, x, ky, kw, Inches(0.95), fill=SOFT, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
            box(s, x, ky + Inches(0.18), Inches(0.06), Inches(0.6), fill=k.get('color', self.accent))
            text(s, x + Inches(0.25), ky + Inches(0.1), kw - Inches(0.3), Inches(0.5), str(k.get('value')),
                 size=24, bold=True, color=k.get('color', INK) if k.get('tint') else INK)
            text(s, x + Inches(0.25), ky + Inches(0.56), kw - Inches(0.3), Inches(0.3), k.get('label', ''),
                 size=10, color=MUTED)
        top = Inches(2.62)
        cols = [('Building / Zone', 0.55, 2.9), ('Stage', 3.5, 1.9), ('Progress', 5.5, 3.1), ('Floors', 8.7, 1.0),
                ('Remarks', 9.75, 3.03)]
        for name, x, w in cols:
            text(s, Inches(x), top, Inches(w), Inches(0.3), name.upper(), size=9, bold=True, color=MUTED)
        box(s, Inches(0.55), top + Inches(0.32), SW - Inches(1.1), Emu(12700), fill='CBD5E1')
        rh = Inches(0.42)
        y = top + Inches(0.4)
        for b in sl.get('buildings', []):
            crit = b.get('critical')
            if crit:
                box(s, Inches(0.55), y, SW - Inches(1.1), rh - Inches(0.04), fill='FEF2F2')
                box(s, Inches(0.55), y, Inches(0.05), rh - Inches(0.04), fill='DC2626')
            name_runs = [(b.get('name', ''), {'bold': True, 'size': 12})]
            if crit:
                name_runs.append(('  CRITICAL', {'bold': True, 'size': 8, 'color': 'DC2626'}))
            text(s, Inches(0.7), y, Inches(2.8), rh, [name_runs], anchor=MSO_ANCHOR.MIDDLE)
            pill = box(s, Inches(3.5), y + Inches(0.07), Inches(1.75), rh - Inches(0.18), fill=b.get('color', MUTED),
                       shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            tf = pill.text_frame
            tf.margin_left = tf.margin_right = Inches(0.05)
            tf.margin_top = tf.margin_bottom = 0
            tf.text = b.get('stage', '')
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            for r in tf.paragraphs[0].runs:
                r.font.size, r.font.bold, r.font.name = Pt(9), True, FONT
                r.font.color.rgb = _rgb(WHITE)
            prog = max(0, min(100, float(b.get('progress') or 0)))
            bw = Inches(2.3)
            box(s, Inches(5.5), y + Inches(0.16), bw, Inches(0.1), fill='E2E8F0', shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            if prog > 0:
                box(s, Inches(5.5), y + Inches(0.16), max(Inches(0.1), int(bw * prog / 100)), Inches(0.1),
                    fill='DC2626' if crit else b.get('color', self.accent), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
            text(s, Inches(7.9), y, Inches(0.75), rh, f'{prog:.0f}%', size=11, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            text(s, Inches(8.7), y, Inches(1.0), rh, b.get('floors') or '—', size=10, color='334155', anchor=MSO_ANCHOR.MIDDLE)
            text(s, Inches(9.75), y, Inches(3.03), rh, b.get('remarks') or '', size=9, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
            box(s, Inches(0.55), y + rh - Inches(0.02), SW - Inches(1.1), Emu(6350), fill=LINE)
            y += rh
        self.footer(s)

    def _table(self, s, headers, widths, rows, top=Inches(1.45), cell_style=None):
        n = len(rows) + 1
        gt = s.shapes.add_table(n, len(headers), Inches(0.55), top, SW - Inches(1.1), Inches(0.4) * n)
        tbl = gt.table
        tblPr = tbl._tbl.tblPr
        style = tblPr.find(qn('a:tableStyleId'))
        if style is None:
            style = etree.SubElement(tblPr, qn('a:tableStyleId'))
        style.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'  # "No Style, Table Grid"
        for i, w in enumerate(widths):
            tbl.columns[i].width = Inches(w)

        def fmt(cell, value, size=10, bold=False, color=INK, fill=None):
            cell.text = str(value if value is not None else '')
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
                    r.font.color.rgb = _rgb(color)
            if fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(fill)
            else:
                cell.fill.background()

        for c, h in enumerate(headers):
            fmt(tbl.cell(0, c), h.upper(), size=9, bold=True, color=WHITE, fill='1E293B')
        for r, row in enumerate(rows, start=1):
            for c, val in enumerate(row):
                st = (cell_style(r - 1, c) if cell_style else None) or {}
                fmt(tbl.cell(r, c), val, **({'fill': 'F8FAFC' if r % 2 == 0 else WHITE} | st))
        return tbl

    def concerns(self, sl):
        s = self.new_slide()
        page = f" ({sl['page']}/{sl['pages']})" if sl.get('pages', 1) > 1 else ''
        self.header(s, (sl.get('title') or 'Key Concerns') + page, 'Issues observed during the drone survey that need attention')
        items = sl.get('items', [])
        rows = [[c.get('label', ''), c.get('title', '') + (('\n' + c['detail']) if c.get('detail') else ''),
                 c.get('building') or '—', (c.get('severity') or 'high').upper(), c.get('action') or '—',
                 c.get('owner') or '—', c.get('due') or '—'] for c in items]

        def style(r, c):
            if c == 3:
                return {'bold': True, 'color': WHITE, 'fill': SEVERITY.get(items[r].get('severity'), 'EF4444'), 'size': 9}
            if c == 0:
                return {'bold': True}
            return None

        self._table(s, ['#', 'Concern', 'Building', 'Severity', 'Action required', 'Responsibility', 'Target'],
                    [0.5, 4.0, 1.6, 1.0, 2.6, 1.3, 1.233], rows, cell_style=style)
        self.footer(s)

    def milestones(self, sl):
        s = self.new_slide()
        page = f" ({sl['page']}/{sl['pages']})" if sl.get('pages', 1) > 1 else ''
        self.header(s, (sl.get('title') or 'Milestones') + page, 'Planned vs. achieved project milestones')
        items = sl.get('items', [])
        rows = [[m.get('title', ''), m.get('building') or '—', m.get('date') or '—', m.get('statusLabel', ''),
                 m.get('note') or ''] for m in items]

        def style(r, c):
            if c == 3:
                return {'bold': True, 'color': WHITE, 'fill': items[r].get('color', MUTED), 'size': 9}
            if c == 0:
                return {'bold': True}
            return None

        self._table(s, ['Milestone', 'Building', 'Date', 'Status', 'Note'], [4.2, 1.9, 1.6, 1.5, 3.033], rows,
                    cell_style=style)
        self.footer(s)

    def end(self, sl):
        s = self.new_slide()
        box(s, 0, 0, SW, SH, fill='0B1220')
        box(s, Inches(0.8), Inches(3.0), Inches(0.09), Inches(1.5), fill=self.accent)
        text(s, Inches(1.1), Inches(2.85), Inches(11), Inches(1.0), sl.get('title') or 'Thank You', size=48, bold=True,
             color=WHITE)
        sub = sl.get('subtitle') or ' · '.join(x for x in [self.info.get('name'), self.info.get('company')] if x)
        text(s, Inches(1.1), Inches(3.9), Inches(11), Inches(0.6), sub, size=18, color='CBD5E1')
        if self.info.get('preparedBy'):
            text(s, Inches(1.1), Inches(6.5), Inches(11), Inches(0.4), 'Prepared by ' + self.info['preparedBy'], size=12,
                 color='94A3B8')

    def build(self, out_path):
        for sl in self.spec.get('slides', []):
            fn = getattr(self, sl.get('type', ''), None)
            if fn and not sl.get('type', '').startswith('_'):
                fn(sl)
        os.makedirs(os.path.dirname(out_path), exist_ok=True)
        self.prs.save(out_path)
        return out_path


def build_pptx(spec, img_dir, out_path):
    return Deck(spec, img_dir).build(out_path)
